from pathlib import Path

import docx
import pytest
from uzru_parser import BlockType, Parser, chunk, parse


@pytest.fixture
def contract(tmp_path: Path) -> Path:
    document = docx.Document()
    document.core_properties.title = "Договор"
    document.add_heading("1. ОБЩИЕ ПОЛОЖЕНИЯ", level=1)
    document.add_paragraph("Настоящий договор является основанием для оказания услуг.")
    document.add_heading("1.1. Основные понятия", level=2)
    document.add_paragraph("первый пункт", style="List Bullet")
    document.add_paragraph("второй пункт", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [("Услуга", "Цена"), ("Хлеб", "100")], strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    document.add_paragraph("Текст после таблицы.")
    path = tmp_path / "contract.docx"
    document.save(path)
    return path


def test_structure_from_styles(contract: Path) -> None:
    result = parse(contract)
    assert [b.type for b in result.blocks] == [
        BlockType.HEADING,
        BlockType.PARAGRAPH,
        BlockType.HEADING,
        BlockType.LIST,
        BlockType.TABLE,
        BlockType.PARAGRAPH,
    ]
    assert [b.level for b in result.blocks if b.type is BlockType.HEADING] == [1, 2]
    assert result.blocks[3].extra["items"] == ["• первый пункт", "• второй пункт"]
    assert result.blocks[4].extra["rows"] == [["Услуга", "Цена"], ["Хлеб", "100"]]


def test_metadata_and_language(contract: Path) -> None:
    result = Parser().parse(contract)
    assert (result.metadata.format, result.metadata.title) == ("docx", "Договор")
    assert result.language.language == "ru"
    assert result.document_id


def test_chunking_keeps_heading_path(contract: Path) -> None:
    chunks = chunk(parse(contract))
    assert chunks[-1].heading_path == ["1. ОБЩИЕ ПОЛОЖЕНИЯ", "1.1. Основные понятия"]
    assert "Услуга | Цена" in chunks[-1].text


def test_uzbek_text_normalized(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_heading("1. UMUMIY QOIDALAR", level=1)
    document.add_paragraph("O'zbekiston Respublikasi ma'lumot va xizmat uchun qoidalar.")
    document.add_heading("1. УМУМИЙ ҚОИДАЛАР", level=1)
    document.add_paragraph("Ўзбекистон Республикаси ўқувчилар учун ғамхўрлик ва маълумот.")
    path = tmp_path / "uz.docx"
    document.save(path)

    blocks = parse(path).blocks
    assert blocks[1].text.startswith("Oʻzbekiston") and "maʼlumot" in blocks[1].text
    assert blocks[1].language.locale == "uz-Latn"
    assert blocks[3].language.locale == "uz-Cyrl"


def test_unstyled_headings_use_heuristics(tmp_path: Path) -> None:
    document = docx.Document()
    for title in ("1. ОБЩИЕ ПОЛОЖЕНИЯ", "2. ПРЕДМЕТ ДОГОВОРА"):
        document.add_paragraph().add_run(title).bold = True
        document.add_paragraph("Текст раздела договора.")
    path = tmp_path / "plain.docx"
    document.save(path)
    assert [b.type for b in parse(path).blocks] == [
        BlockType.HEADING,
        BlockType.PARAGRAPH,
        BlockType.HEADING,
        BlockType.PARAGRAPH,
    ]


def test_lone_bold_numbered_line_in_docx_is_a_list_item(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_paragraph("Talablar quyidagilardan iborat.")
    document.add_paragraph().add_run("4. Hujjatlarni oʻz vaqtida topshirish").bold = True
    path = tmp_path / "item.docx"
    document.save(path)
    assert [b.type for b in parse(path).blocks] == [BlockType.PARAGRAPH, BlockType.LIST]


def test_explicit_page_breaks_advance_pages(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_paragraph("Первая страница.")
    document.add_page_break()
    document.add_paragraph("Вторая страница.")
    path = tmp_path / "pages.docx"
    document.save(path)

    result = parse(path)
    assert [b.page for b in result.blocks] == [1, 2]
    assert result.metadata.page_count == 2
    assert result.metadata.extra["pages_approximate"]


def test_docx_mixed_alphabet_words_are_fixed(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_paragraph("Vаlyutа bozori va mа’lumot tizimi.")
    path = tmp_path / "mixed.docx"
    document.save(path)
    result = parse(path)
    assert result.blocks[0].text == "Valyuta bozori va maʼlumot tizimi."
    assert result.metadata.extra["mixed_script_words_fixed"] == 2


def bold_paragraph(document: docx.document.Document, text: str) -> None:
    document.add_paragraph().add_run(text).bold = True


def test_sentence_broken_over_paragraphs_by_ocr_is_joined(tmp_path: Path) -> None:
    document = docx.Document()
    bold_paragraph(document, "І. Общие положения")
    document.add_paragraph("Участниками системы обеспечения информационной")
    document.add_paragraph("безопасности являются: собственники объектов инфраструктуры.")
    document.add_paragraph("В настоящей Доктрине используются следующие основные")
    bold_paragraph(document, "ПОНЯТИЯ:")
    document.add_paragraph("а) национальные интересы в информационной сфере;")
    bold_paragraph(document, "ІІ. Национальные интересы")
    document.add_paragraph("Национальными интересами являются права граждан.")
    path = tmp_path / "ocr.docx"
    document.save(path)

    blocks = parse(path).blocks
    assert [b.text for b in blocks if b.type is BlockType.HEADING] == [
        "І. Общие положения",
        "ІІ. Национальные интересы",
    ]
    assert [(b.type, b.text) for b in blocks[:3]] == [
        (BlockType.HEADING, "І. Общие положения"),
        (
            BlockType.PARAGRAPH,
            "Участниками системы обеспечения информационной безопасности являются: "
            "собственники объектов инфраструктуры.",
        ),
        (BlockType.PARAGRAPH, "В настоящей Доктрине используются следующие основные ПОНЯТИЯ:"),
    ]


def test_finished_paragraphs_in_docx_are_not_joined(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_paragraph("(107-modda Qonuniga asosan chiqarilgan — 14.02.2025-y., 0144-son)")
    document.add_paragraph("shaxsiy jamgʻarib boriladigan pensiya hisobvaragʻi raqami.")
    document.add_paragraph("Birinchi gap tugadi.")
    document.add_paragraph("keyingi qator kichik harf bilan.")
    document.add_paragraph("Matn shu yerda tugadi.")
    bold_paragraph(document, "UMUMIY QOIDALAR")
    document.add_paragraph("Ushbu qoidalar xizmat tartibini belgilaydi.")
    path = tmp_path / "laws.docx"
    document.save(path)

    blocks = parse(path).blocks
    assert len(blocks) == 7
    assert blocks[5].type is BlockType.HEADING


def save(document: docx.document.Document, path: Path) -> Path:
    document.save(str(path))
    return path


def centred_bold(document: docx.document.Document, text: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    paragraph = document.add_paragraph()
    paragraph.add_run(text).bold = True
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def guide(tmp_path: Path) -> Path:
    document = docx.Document()
    centred_bold(document, "Foydalanuvchi qoʻllanmasi")
    centred_bold(document, "Markaziy bankning mobil ilovasi")
    bold_paragraph(document, "1. Umumiy maʼlumot")
    document.add_paragraph("Ilova foydalanuvchilarga bank tizimi haqida maʼlumot beradi.")
    bold_paragraph(document, "Kredit limiti")
    document.add_paragraph("Kredit limiti bank berishga tayyor boʻlgan eng katta summadir.")
    paragraph = document.add_paragraph()
    paragraph.add_run("Overdraft.").bold = True
    paragraph.add_run(" Overdraft hisobdagidan koʻp mablagʻdan foydalanish imkoniyatidir.")
    bold_paragraph(document, "Ilovadan qanday foydalanaman?")
    document.add_paragraph("Ilovani yuklab oling va roʻyxatdan oʻting.")
    for step in ("Ilovani oching", "Raqamni kiriting", "Kodni tasdiqlang"):
        document.add_paragraph(step, style="List Number")
    bold_paragraph(document, "2. Xavfsizlik")
    document.add_paragraph("Parolni hech kimga bermang.")
    return save(document, tmp_path / "guide.docx")


def test_bold_standalone_paragraphs_are_subheadings(tmp_path: Path) -> None:
    blocks = parse(guide(tmp_path)).blocks
    headings = [(b.text, b.level) for b in blocks if b.type is BlockType.HEADING]
    assert headings == [
        ("1. Umumiy maʼlumot", 1),
        ("Kredit limiti", 2),
        ("Ilovadan qanday foydalanaman?", 2),
        ("2. Xavfsizlik", 1),
    ]
    run_in = next(b for b in blocks if b.text.startswith("Overdraft."))
    assert run_in.type is BlockType.PARAGRAPH


def test_word_numbered_list_keeps_its_numbers(tmp_path: Path) -> None:
    lists = [b for b in parse(guide(tmp_path)).blocks if b.type is BlockType.LIST]
    assert lists[-1].extra["items"] == [
        "1. Ilovani oching",
        "2. Raqamni kiriting",
        "3. Kodni tasdiqlang",
    ]


def test_leading_centred_bold_lines_are_the_title(tmp_path: Path) -> None:
    document = parse(guide(tmp_path))
    assert document.metadata.title == "Foydalanuvchi qoʻllanmasi Markaziy bankning mobil ilovasi"
    chunks = chunk(document)
    assert not any(c.text.startswith("Foydalanuvchi qoʻllanmasi") for c in chunks)
    assert chunks[0].heading_path[0] == "1. Umumiy maʼlumot"


def test_bullet_list_stays_bulleted(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_paragraph("Ilova quyidagilarni taqdim etadi:")
    for item in ("maʼlumot", "xizmatlar"):
        document.add_paragraph(item, style="List Bullet")
    lists = [
        b for b in parse(save(document, tmp_path / "b.docx")).blocks if b.type is BlockType.LIST
    ]
    assert lists[0].extra["items"] == ["• maʼlumot", "• xizmatlar"]


def test_corrupt_or_empty_docx_gives_a_clear_error(tmp_path: Path) -> None:
    from uzru_parser import DocumentError

    corrupt = tmp_path / "corrupt.docx"
    corrupt.write_bytes(b"PK\x03\x04 not really a zip")
    empty = tmp_path / "empty.docx"
    empty.write_bytes(b"")
    with pytest.raises(DocumentError, match="not a valid DOCX"):
        parse(corrupt)
    with pytest.raises(DocumentError, match="empty"):
        parse(empty)
    with pytest.raises(DocumentError, match="Unsupported"):
        parse(tmp_path / "notes.txt")


def centred(document: docx.document.Document, text: str, bold: bool = False) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    paragraph = document.add_paragraph()
    paragraph.add_run(text).bold = bold
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def legal_act(tmp_path: Path) -> Path:
    """An act as published: a byte-order mark, the issuer and kind of act in plain centred
    lines, the bold title, chapters, a signature, then an appendix whose stamp ends with its
    label on a line of its own, the appendix title and its chapters."""
    document = docx.Document()
    document.add_paragraph("﻿")
    centred(document, "OʻZBEKISTON RESPUBLIKASI VAZIRLAR MAHKAMASINING")
    centred(document, "QARORI")
    centred(document, "Davlat xaridlari tizimini takomillashtirish toʻgʻrisida", bold=True)
    document.add_paragraph("Vazirlar Mahkamasi qaror qiladi:")
    centred(document, "I. Umumiy qoidalar", bold=True)
    document.add_paragraph("1. 2026-yil 1-yanvardan boshlab yangi tartib joriy etilsin.")
    centred(document, "II. Yakuniy qoidalar", bold=True)
    document.add_paragraph("2. Ushbu qaror rasmiy eʼlon qilingan kundan kuchga kiradi.")
    bold_paragraph(document, "Oʻzbekiston Respublikasining Bosh vaziri A. ARIPOV")
    centred(document, "Toshkent sh.,")
    paragraph = document.add_paragraph()
    paragraph.add_run("Vazirlar Mahkamasining 2026-yil 8-sentabrdagi 480-son qaroriga")
    paragraph.add_run().add_break()
    paragraph.add_run("1-ILOVA")
    centred(document, "Xaridlar narxini aniqlash tartibi toʻgʻrisida nizom", bold=True)
    centred(document, "1-bob. Umumiy qoidalar", bold=True)
    bold_paragraph(document, "Nizomning asosiy tushunchalari quyidagilardan iborat:")
    document.add_paragraph("narx – tovarning bozordagi qiymati.")
    centred(document, "2-bob. Yakuniy qoidalar", bold=True)
    document.add_paragraph("Nizom talablari barcha buyurtmachilarga tatbiq etiladi.")
    return save(document, tmp_path / "act.docx")


def test_a_legal_act_has_its_title_as_metadata_and_a_clean_outline(tmp_path: Path) -> None:
    result = parse(legal_act(tmp_path))
    assert result.metadata.title == "Davlat xaridlari tizimini takomillashtirish toʻgʻrisida"
    headings = [(b.level, b.text) for b in result.blocks if b.type is BlockType.HEADING]
    texts = [text for _, text in headings]
    # The signature and the ":" lead-in are no headings; the title is metadata.
    assert not any("ARIPOV" in text or text.endswith(":") for text in texts)
    assert "Davlat xaridlari tizimini takomillashtirish toʻgʻrisida" not in texts
    assert texts[:3] == ["I. Umumiy qoidalar", "II. Yakuniy qoidalar", "1-ILOVA"]
    chunks = chunk(result)
    path = next(c.heading_path for c in chunks if "tatbiq etiladi" in c.text)
    assert path[0] == "1-ILOVA" and path[-1] == "2-bob. Yakuniy qoidalar"


def test_cells_merged_down_are_read_once(tmp_path: Path) -> None:
    document = docx.Document()
    table = document.add_table(rows=4, cols=3)
    for cell, text in zip(
        table.rows[0].cells, ("Bosqichlar", "Tadbirlar", "Muddatlar"), strict=True
    ):
        cell.text = text
    for column in range(3):  # the header cells cover two rows
        table.cell(0, column).merge(table.cell(1, column))
    for cell, text in zip(
        table.rows[2].cells, ("1-bosqich", "Reja tuziladi", "1-aprel"), strict=True
    ):
        cell.text = text
    table.cell(2, 0).merge(table.cell(3, 0))
    table.cell(3, 1).text = "Reja tasdiqlanadi"
    table.cell(3, 2).text = "1-may"
    rows = parse(save(document, tmp_path / "merged.docx")).blocks[0].extra["rows"]
    assert rows == [
        ["Bosqichlar", "Tadbirlar", "Muddatlar"],
        ["1-bosqich", "Reja tuziladi", "1-aprel"],
        ["", "Reja tasdiqlanadi", "1-may"],
    ]


def test_an_english_document_has_no_uzbek_language(tmp_path: Path) -> None:
    document = docx.Document()
    for _ in range(6):
        document.add_paragraph(
            "Python is a programming language that lets you work quickly and integrate "
            "systems more effectively, and this guide explains how to write it well."
        )
    document.add_paragraph("Bu kitob dasturlash tilini oʻrganish uchun moʻljallangan.")
    result = parse(save(document, tmp_path / "english.docx"))
    assert result.language.language == "unknown"
    assert all(c.language != "uz" or "kitob" in c.text for c in chunk(result))


def test_a_placeholder_title_property_gives_way_to_the_documents_own(tmp_path: Path) -> None:
    document = docx.Document()
    document.core_properties.title = "Untitled"
    centred_bold(document, "Introducing DuoWrite")
    document.add_paragraph("DuoWrite is a pen that is also a pencil, made for everyday writing.")
    assert parse(save(document, tmp_path / "duo.docx")).metadata.title == "Introducing DuoWrite"
    empty = docx.Document()
    empty.core_properties.title = "Untitled"
    empty.add_paragraph("Plain text without a title of its own.")
    assert parse(save(empty, tmp_path / "plain.docx")).metadata.title is None


def test_a_bold_line_like_an_outline_title_is_its_sibling(tmp_path: Path) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    def title(text: str, outline: bool) -> None:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(text)
        run.bold, run.font.size = True, Pt(13.5)
        if outline:
            level = OxmlElement("w:outlineLvl")
            level.set(qn("w:val"), "2")
            paragraph._p.get_or_add_pPr().append(level)

    document = docx.Document()
    title("Настраиваемый конструктор правил опозданий", outline=True)
    document.add_paragraph("HR задаёт допустимое опоздание и правила для каждой смены.")
    title("Автоматическая классификация статусов", outline=False)
    document.add_paragraph("Система сама определяет статус каждого сотрудника за день.")
    title("Уведомления и эскалации", outline=True)
    document.add_paragraph("Руководитель получает уведомление о каждом нарушении.")
    blocks = parse(save(document, tmp_path / "rules.docx")).blocks
    levels = [b.level for b in blocks if b.type is BlockType.HEADING]
    assert levels == [1, 1, 1]


def test_text_in_text_boxes_is_read_once(tmp_path: Path) -> None:
    from docx.oxml import parse_xml

    document = docx.Document()
    document.add_paragraph("Why DuoWrite stands out")
    w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    mc = 'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
    text = "<w:p><w:r><w:t>Two-in-one solution for every desk</w:t></w:r></w:p>"
    box = f"<w:txbxContent>{text}</w:txbxContent>"
    run = parse_xml(
        f'<w:r {w} {mc}><mc:AlternateContent><mc:Choice Requires="wps"><w:pict>{box}</w:pict>'
        f"</mc:Choice><mc:Fallback><w:pict>{box}</w:pict></mc:Fallback></mc:AlternateContent></w:r>"
    )
    document.paragraphs[0]._p.append(run)
    read = " ".join(b.text for b in parse(save(document, tmp_path / "boxes.docx")).blocks)
    assert read.count("Two-in-one solution for every desk") == 1
