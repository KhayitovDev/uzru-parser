"""Page confidence: each check lowers it and names its issue; clean pages stay at 1."""

from pathlib import Path

from docx import Document as WordDocument
from helpers import make_pdf
from uzru_parser import parse
from uzru_parser.confidence import page_confidence
from uzru_parser.paragraphs import Line
from uzru_parser.profile import StyleProfile, build_profile
from uzru_parser.structure import RawBlock

WORDS = (
    "Markaziy bank pul muomalasini tartibga soladi va tijorat banklari faoliyatini nazorat "
    "qiladi hamda davlat iqtisodiyotining barqarorligini taʼminlaydi"
).split()


def text_line(index: int, x0: float = 72, x1: float = 520, text: str | None = None) -> Line:
    words = text if text is not None else " ".join(WORDS[index % 5 : index % 5 + 9])
    return Line(
        text=words, bbox=(x0, 100 + 14 * index, x1, 112 + 14 * index), size=11.0, chars=len(words)
    )


def body_page() -> list[Line]:
    return [text_line(i) for i in range(30)]


def test_a_clean_page_is_fully_confident() -> None:
    assert page_confidence([body_page()], [], [], None) == (1.0, [])


def test_columns_read_as_one_region_cost_confidence() -> None:
    lines = [text_line(i, 72, 280) for i in range(20)] + [text_line(i, 310, 520) for i in range(20)]
    score, issues = page_confidence([lines], [], [], None)
    assert issues == ["columns"] and score < 0.5
    assert page_confidence([lines[:20], lines[20:]], [], [], None) == (1.0, [])


def test_a_page_of_short_fragments_is_unsure() -> None:
    lines = [text_line(i, text=t) for i, t in enumerate(["=", "p&", "G &", "0", "x2", "(18)"] * 4)]
    score, issues = page_confidence([lines + body_page()[:8]], [], [], None)
    assert "short_lines" in issues and score < 0.7


def test_broken_text_layers_are_found() -> None:
    garbled = "Ïðîôåññèîíàëüíûé ñòàíäàðò óòâåðæäåí ïðèêàçîì ìèíèñòåðñòâà òðóäà"
    lines = [text_line(i, text=garbled) for i in range(10)]
    assert page_confidence([lines], [], [], None) == (0.0, ["text_layer"])
    marked = [text_line(0, text="Bank �� hisoboti"), *body_page()]
    assert page_confidence([marked], [], [], None) == (0.0, ["text_layer"])


def test_poor_tables_cost_confidence() -> None:
    poor = [["Koʻrsatkich", "", ""], ["", "", ""], ["Daromad (birlash", ")", ""], ["", "", ""]]
    good = [["Koʻrsatkich", "2023"], ["Daromad", "100"], ["Xarajat", "80"]]
    score, issues = page_confidence([body_page()], [poor], [], None)
    assert issues == ["tables"] and score < 0.5
    assert page_confidence([body_page()], [good], [], None) == (1.0, [])


def block(text: str, y: float, size: float = 11.0, bold: float = 0.0) -> RawBlock:
    return RawBlock(text=text, page=1, bbox=(72, y, 520, y + 12), font_size=size, bold=bold)


def test_prominent_blocks_outside_the_heading_styles_conflict() -> None:
    body = [block(" ".join(WORDS), 100 + 40 * i) for i in range(6)]
    odd = [block(f"Alohida qalin satr {i}", 400 + 30 * i, size=13, bold=1.0) for i in range(3)]
    profile = build_profile(body, {1: 800.0})
    assert profile.headings == {}
    score, issues = page_confidence([body_page()], [], [*body, *odd], profile)
    assert issues == ["headings"] and score == 0.0
    empty = StyleProfile()
    assert page_confidence([body_page()], [], [*body, *odd], empty) == (1.0, [])


def test_pdf_pages_carry_their_confidence(tmp_path: Path) -> None:
    text = " ".join(WORDS) + ". " + " ".join(reversed(WORDS)) + "."
    document = parse(make_pdf(tmp_path / "doc.pdf", [text, text]))
    assert [page.extra for page in document.pages] == [{"confidence": 1.0, "issues": []}] * 2
    assert all("extra" in page for page in document.to_dict()["pages"])


def test_docx_pages_have_no_extra_key(tmp_path: Path) -> None:
    word = WordDocument()
    word.add_paragraph(" ".join(WORDS))
    path = tmp_path / "doc.docx"
    word.save(str(path))
    pages = parse(path).to_dict()["pages"]
    assert pages and all("extra" not in page for page in pages)
